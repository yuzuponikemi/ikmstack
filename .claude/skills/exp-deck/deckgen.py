"""exp-deck レンダラ —— 報告デッキ(.pptx)生成の「仕組み」.

実験側は中身(deck.py = 純データ)だけを持ち、生成の仕組みはこのファイルに集約する。
deck.py のスキーマは下部のドキュメント文字列と SKILL.md を参照。

デッキは「伝達(コンサル的メッセージ管理)」と「議論の土台(研究者のピアレビュー)」を両立させる:
型 = 背景→結論(条件付き)→行動→根拠(出典+トレードオフ)→確度(事実/推論・N/誤差)→Discussion & Future Work。
詳細な原則(技術スライド5大原則)は SKILL.md「資料作りの根幹」を参照。

実行:
  uv run --with python-pptx .claude/skills/exp-deck/deckgen.py experiments/<ID>/deck.py

deck.py(content)が定義すべきモジュール属性:
  OUT_NAME : str           出力の基底名(拡張子は無視。生成時に _YYYYMMDD-HHMM が付与される)
  TAG      : str           フッターに出す実験ID(例 "E058")
  TITLE    : str           表紙タイトル(改行可)
  SUBTITLE : str           表紙サブタイトル(問い)
  META     : str           表紙メタ行(範囲・日付)
  SLIDES   : list[dict]    スライド定義(下記)

SLIDES の各要素:
  {"type": "title"}
      表紙。TITLE/SUBTITLE/META を使う。
  {"type": "bullets", "title": str, "color": "navy"|"red"|"gray",
   "items": [(level, text[, color, bold]), ...], "size": int, "gap": int}
      全幅の箇条書き。level 0=■見出し / 1=・子。color/bold は項目ごとに省略可。
  {"type": "figure", "title": str, "image": "<figures内のファイル名>",
   "caption": str, "items": [...](あれば2カラム), "image_side": "left"|"right",
   "image_w": float, "image_h": float, "left": float, "source": str}
      図スライド。items があれば「図+箇条書き」の2カラム(image_side で図の左右、既定 left)。
      items 無しなら図を主役に中央寄せ。縦長/横長の図は image_side="right" 側が収まりやすい。

  "source": str (図/グラフ/表を載せるスライドでは必須)
      スライド下部に出る出典・来歴。**そのスライドだけ見て読者が理解できる**ことが基準
      (長くなってよい。engine が複数行表示し、図を自動で縮めて場所を空ける)。書くべきこと:
      ①データの素性 — **加工済み中間データなら、それが何か・どの一次データから来たかを平易に説明**
        (例「washparallel_runs.csv = 各装置が運用中に残す生の流体ログ(FluidSysLog_*)から抵抗チェック
         1回=1行に整理した中間データ」)。**一次ログ(Drive の FluidSysLog 等)そのものなら引用のみで可**。
      ②処理 — スクリプトが何をするかを**日本語で**(ファイル名だけでなく動作。例「装置別に集計し robust CV を算出」)。
      ③可視化 — 何を軸/色/棒/線にしたか。④意図 — その図で何を示したいか。
      最後に "(生成: scripts/foo.py)" を添えて再現性も残す。
      NG例(読者が辿れない): "scripts/foo.py が bar.csv を集計"。
      bullets スライドでも表を載せるなら付ける。title スライドや純粋な文章スライドには不要。

色は文字列で指定。テーマは組織ブランド(単一ソース = brand/tokens.py):
  "navy"/"purple"=テーマメインカラー(見出し) / "green"=アクセントカラー / "red"=制約・トレードオフ・要注意 /
  "gray"=補足 / "ink"=本文濃色。表紙はカラーバンド+背景画像+ロゴ、コンテンツは白地+細罫線、
  フッターにロゴ+ページ番号(図の出典スライドではロゴを省き出典帯を優先)。
  HTML/Artifact 用の同等スタイルは brand/brand.css(tokens.py が生成)。実験側に pptx 依存を持ち込ませない。
"""

import importlib.util
import sys
from datetime import datetime
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

# ---- テーマ: 組織ブランド(単一ソース = brand/tokens.py) ----------------
sys.path.insert(0, str(Path(__file__).parent / "brand"))
import tokens as B  # noqa: E402


def _rgb(h):
    return RGBColor(int(h[0:2], 16), int(h[2:4], 16), int(h[4:6], 16))


PURPLE = _rgb(B.PURPLE)
PURPLE_HEAD = _rgb(B.PURPLE_HEAD)
GREEN = _rgb(B.GREEN)
INK = _rgb(B.INK)
GRAY = _rgb(B.GRAY)
FOOTER_BG = _rgb(B.FOOTER_BG)
RED = _rgb(B.RED)
WHITE = _rgb(B.WHITE)
SUB1 = _rgb("D9C6E6")  # 表紙バンド上のサブタイトル(淡ラベンダー)
NAVY = PURPLE_HEAD      # 後方互換エイリアス
FONT = B.FONT_JP
COLORS = {k: _rgb(v) for k, v in B.SEMANTIC.items()}

ASSETS = Path(__file__).parent / "brand" / "assets"
LOGO = ASSETS / "logo.png"
BG_IMG = ASSETS / "background.png"

_asset_warned: set = set()


def _usable(path: Path) -> bool:
    """ブランド画像が **実際に読める** か。`exists()` だけでは足りない。

    意匠画像は差し替え前提のプレースホルダなので、壊れていても**デッキ本体の生成は
    続行すべき**(ロゴは装飾で、成果物はスライド)。ここを `exists()` だけで守っていたため、
    リポジトリ初版から入っていた破損 PNG で **exp-deck が表紙生成で必ず落ちていた**
    (PIL.UnidentifiedImageError)。存在と使用可能は別物。
    """
    if not path.exists():
        return False
    try:
        from PIL import Image  # python-pptx の依存として入っている
    except ImportError:
        return True          # 判定できないなら従来どおり任せる
    try:
        with Image.open(path) as im:
            im.load()
        return True
    except Exception as e:                                    # noqa: BLE001
        if path not in _asset_warned:
            _asset_warned.add(path)
            print(f"WARN: ブランド画像を読めないので省略します: {path} ({e})",
                  file=sys.stderr)
        return False


SW, SH = Inches(13.333), Inches(7.5)  # 16:9


def _color(c):
    if c is None:
        return INK
    return COLORS.get(c, INK) if isinstance(c, str) else c


def set_font(run, name=FONT, size=None, bold=None, color=None, italic=None):
    """日本語が崩れないよう latin/ea/cs すべての typeface を設定する。"""
    run.font.name = name
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = rPr.makeelement(qn(tag), {})
            rPr.append(el)
        el.set("typeface", name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.font.bold = bold
    if italic is not None:
        run.font.italic = italic
    if color is not None:
        run.font.color.rgb = _color(color)


def _blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def _rect(slide, left, top, width, height, color):
    r = slide.shapes.add_shape(1, left, top, width, height)  # 1 = rectangle
    r.fill.solid()
    r.fill.fore_color.rgb = color
    r.line.fill.background()
    r.shadow.inherit = False
    return r


def _title(slide, text, color=None, size=26):
    """コンテンツ見出し: 白ヘッダに濃色タイトル + テーマカラーの細罫(標準ブランドスタイル)."""
    box = slide.shapes.add_textbox(Inches(0.5), Inches(0.30), SW - Inches(1.0), Inches(0.85))
    box.text_frame.word_wrap = True
    r = box.text_frame.paragraphs[0].add_run()
    r.text = text
    set_font(r, size=size, bold=True, color=_color(color) if color else INK)
    _rect(slide, Inches(0.5), Inches(1.20), SW - Inches(1.0), Inches(0.035), PURPLE)  # accent rule


def _bullets(slide, items, left, top, width, height, base_size=18, gap=6):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    first = True
    for item in items:
        level, text = item[0], item[1]
        color = item[2] if len(item) > 2 else NAVY
        bold = item[3] if len(item) > 3 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.level = level
        p.space_after = Pt(gap)
        r = p.add_run()
        r.text = ("■ " if level == 0 else "・ ") + text
        set_font(r, size=base_size if level == 0 else base_size - 3,
                 bold=bold or level == 0, color=color)


def _image_fit(slide, path, left, top, max_w, max_h, caption=None):
    pic = slide.shapes.add_picture(str(path), left, top, width=max_w)
    if pic.height > max_h:
        pic._element.getparent().remove(pic._element)
        pic = slide.shapes.add_picture(str(path), left, top, height=max_h)
    pic.left = int(left + (max_w - pic.width) / 2)
    if caption:
        cap = slide.shapes.add_textbox(left, top + pic.height + Inches(0.05), max_w, Inches(0.4))
        cp = cap.text_frame.paragraphs[0]
        cp.alignment = PP_ALIGN.CENTER
        r = cp.add_run()
        r.text = caption
        set_font(r, size=11, italic=True, color=GRAY)


def _footer(slide, n, full=True):
    """標準フッター: 細罫 + 左にロゴ + 右にページ番号。
    full=False(出典スライド)はロゴ/罫を省き、下部の出典帯と衝突させない。"""
    if full:
        _rect(slide, Inches(0.0), Inches(7.05), SW, Pt(1.2), FOOTER_BG)  # hairline
        if _usable(LOGO):
            slide.shapes.add_picture(str(LOGO), Inches(0.45), Inches(7.13), height=Inches(0.25))
    box = slide.shapes.add_textbox(Inches(11.8), Inches(7.08), Inches(1.4), Inches(0.35))
    p = box.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    r = p.add_run()
    r.text = n
    set_font(r, size=10, color=GRAY)


def _source(slide, text):
    """Provenance strip at the bottom of a figure/table slide: what the data is (esp.
    if it is a *derived/processed* dataset, not a raw log), how it was processed &
    visualized, and the intent — written so a reader understands the slide on its own.
    Multi-line is expected; the strip spans the lower band, clear of the page-number
    footer at the right. The figure is auto-shrunk (see _ih) so it clears this band."""
    box = slide.shapes.add_textbox(Inches(0.4), Inches(6.45), Inches(11.3), Inches(1.0))
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    r = p.add_run()
    r.text = "データ: " + text
    set_font(r, size=9, color=GRAY)


def _title_slide(prs, c):
    """表紙: 上=テーマカラーバンド+背景画像、白タイトル、下=META+ロゴ。"""
    s = _blank(prs)
    BAND_H = Inches(4.05)
    _rect(s, 0, 0, SW, BAND_H, PURPLE)
    if _usable(BG_IMG):
        # 幅 6.13" は「標準ブランド画像なら高さ≈BAND_H になる」値。差し替え画像の
        # 縦横比が違うとバンドを越えて META 行に被るので、越えたら高さで入れ直す。
        iw = Inches(6.13)
        pic = s.shapes.add_picture(str(BG_IMG), int(SW - iw), 0, width=iw)
        if pic.height > BAND_H:
            pic._element.getparent().remove(pic._element)
            pic = s.shapes.add_picture(str(BG_IMG), int(SW - iw), 0, height=BAND_H)
        pic.left = int(SW - pic.width); pic.top = 0
    # タイトル(白)はバンド左、虹彩(left≈7.2")に被らない幅に
    tb = s.shapes.add_textbox(Inches(0.8), Inches(1.05), Inches(6.1), Inches(2.6))
    tf = tb.text_frame; tf.word_wrap = True
    r = tf.paragraphs[0].add_run(); r.text = getattr(c, "TITLE", "")
    set_font(r, size=33, bold=True, color=WHITE)
    p2 = tf.add_paragraph(); p2.space_before = Pt(14)
    r2 = p2.add_run(); r2.text = getattr(c, "SUBTITLE", "")
    set_font(r2, size=16, color=SUB1)
    # META は白地に
    mb = s.shapes.add_textbox(Inches(0.8), Inches(4.45), SW - Inches(1.6), Inches(1.0))
    mtf = mb.text_frame; mtf.word_wrap = True
    rm = mtf.paragraphs[0].add_run(); rm.text = getattr(c, "META", "")
    set_font(rm, size=14, color=INK)
    if _usable(LOGO):
        s.shapes.add_picture(str(LOGO), Inches(0.8), Inches(6.82), height=Inches(0.32))


def _figure_roots(exp_dir: Path) -> list[str]:
    """図の探索先を .lab-config.json の `allowed_figure_roots` から取る(既定 `figures`)。

    ノート根は experiments/<ID>/ の 2 つ上。見つからなければ既定にフォールバックする。
    """
    roots = None
    for cand in (exp_dir.parent.parent, exp_dir.parent):
        cfg = cand / ".lab-config.json"
        if cfg.exists():
            try:
                import json
                roots = json.loads(cfg.read_text(encoding="utf-8")).get("allowed_figure_roots")
            except (OSError, ValueError):
                roots = None
            break
    roots = [str(r).strip("/ ") for r in (roots or []) if str(r).strip()]
    return roots or ["figures"]


def _resolve_image(name: str, fig_dir: Path) -> Path:
    """deck.py の `image` を実ファイルへ解決する。

    `figures/` 決め打ちだと、図の出力先が `_generated/` のノートで必ず落ちる
    (実験側が deck.py に絶対パスを書く回避策に追い込まれる — 中身は純データに保ちたい)。
    探索順:
      1. 絶対パス / fig_dir 直下(従来の互換動作)
      2. `allowed_figure_roots` の各ルート配下を再帰探索(`plots/` などの入れ子に対応)
    """
    p = Path(name)
    if p.is_absolute():
        return p
    direct = fig_dir / name
    if direct.exists():
        return direct
    exp_dir = fig_dir.parent
    for root in _figure_roots(exp_dir):
        base = exp_dir / root
        if not base.is_dir():
            continue
        cand = base / name
        if cand.exists():
            return cand
        hits = sorted(base.rglob(Path(name).name))
        if hits:
            return hits[0]
    return direct   # 見つからなければ従来のパスで落として、エラー文言を素直にする


def render(content, fig_dir: Path, out_path: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = SW, SH
    tag = getattr(content, "TAG", "")

    for pos, spec in enumerate(content.SLIDES, start=1):
        kind = spec["type"]
        if kind == "title":
            _title_slide(prs, content)
            continue

        s = _blank(prs)
        _title(s, spec["title"], color=spec.get("color"))
        foot = f"{tag} ・ {pos}"

        if kind == "bullets":
            _bullets(s, spec["items"], Inches(0.7), Inches(1.5),
                     SW - Inches(1.4), Inches(5.4),
                     base_size=spec.get("size", 18), gap=spec.get("gap", 6))

        elif kind == "figure":
            img = _resolve_image(spec["image"], fig_dir)
            # when a source strip is present, cap image height so image+caption clears
            # the (taller, multi-line) provenance band at the slide bottom
            def _ih(default):
                cap = 4.3 if spec.get("source") else default
                return Inches(min(spec.get("image_h", default), cap))
            if spec.get("items"):  # 図+箇条書きの2カラム(image_side で図の左右を選択)
                cap = spec.get("caption")
                size = spec.get("size", 15)
                if spec.get("image_side", "left") == "right":
                    _bullets(s, spec["items"], Inches(0.55), Inches(1.55),
                             Inches(7.2), Inches(5.4), base_size=size, gap=5)
                    _image_fit(s, img, Inches(spec.get("left", 8.0)), Inches(1.6),
                               Inches(spec.get("image_w", 4.9)), _ih(4.9),
                               caption=cap)
                else:
                    _image_fit(s, img, Inches(spec.get("left", 0.4)), Inches(1.45),
                               Inches(spec.get("image_w", 7.4)), _ih(5.4),
                               caption=cap)
                    _bullets(s, spec["items"], Inches(8.1), Inches(1.55),
                             Inches(4.9), Inches(5.4), base_size=size, gap=5)
            else:  # 図を主役に中央寄せ
                _image_fit(s, img, Inches(spec.get("left", 1.8)), Inches(1.5),
                           Inches(spec.get("image_w", 9.7)), _ih(5.4),
                           caption=spec.get("caption"))
        else:
            raise ValueError(f"unknown slide type: {kind!r}")

        if spec.get("source"):
            _source(s, spec["source"])
        _footer(s, foot, full=not bool(spec.get("source")))

    out_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(out_path))
    return len(prs.slides._sldIdLst)


def _load(path: Path):
    spec = importlib.util.spec_from_file_location("deck_content", path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    return mod


def main(argv):
    if len(argv) != 2:
        sys.exit("usage: deckgen.py <experiments/<ID>/deck.py>")
    content_path = Path(argv[1]).resolve()
    if not content_path.exists():
        sys.exit(f"not found: {content_path}")
    exp_dir = content_path.parent
    content = _load(content_path)
    # 生成日時をファイル名に付与 → 再生成で既存(編集済み)ファイルを上書きしない
    stem = Path(getattr(content, "OUT_NAME", "report.pptx")).stem
    ts = datetime.now().strftime("%Y%m%d-%H%M")
    out = exp_dir / "outputs" / f"{stem}_{ts}.pptx"
    n = render(content, exp_dir / "figures", out)
    print(f"saved: {out}  ({out.stat().st_size/1024:.0f} KB, {n} slides)")


if __name__ == "__main__":
    main(sys.argv)
